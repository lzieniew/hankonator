from docx.shared import Pt

from generation import Project

class Stage6(object):

    def __init__(self, erd, transactions):
        self.erd = erd
        self.transactions = transactions

    def build(self, document):
        header = document.add_paragraph()
        header.add_run('6. Transakcje').font.size = Pt(Project.HEADER_SIZE)
        header.add_run().add_break()
